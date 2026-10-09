"""Compare roberta-large alone with the bare LLM predictions and the ground truth on every test set.

For every data/test_data/<dataset>-<seed>.xlsx:
  1. roberta-large predictions: taken from ./outputs/roberta_large/<test file name>.csv, written by
     predict_all_files_roberta_large.py. A pair without that file is trained here first, with the
     same procedure (roberta-large alone, Optuna over the learning-rate x batch-size grid and
     epochs, see that script); --retrain trains every pair again.
  2. LLM predictions: GPT-5 nano, 5-shot, from ./outputs/predictions/test_gpt-5-nano_5_shot/.
  3. Both are compared with each other and with the ground truth, sentence by sentence.

Two colour columns per sentence:
    agreement_color     green  = LLM and roberta-large predict the same class
                        white  = they differ
    correctness_color   blue   = both are correct
                        yellow = only roberta-large is correct
                        green  = only the LLM is correct
                        white  = both are wrong

Writes to ./outputs/roberta_vs_llm/:
    comparison.csv    every test sentence of every pair, colours as text
    comparison.xlsx   the same table with the two colour columns filled in those colours
    summary.csv       per pair and per dataset: share of agreement / only roberta / only LLM / both / neither

    python testing/compare_roberta_llm.py                       # all test sets
    python testing/compare_roberta_llm.py --datasets mm-test    # datasets containing these
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

TESTING_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TESTING_DIR))

import pandas as pd  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

import predict_all_files_roberta_large as roberta_large  # noqa: E402
from predict_all_files_labelFusion import LLM_TEST_DIR, OUT_DIR, TEST_DIR, TEXT_COLUMN, llm_predictions, paired_train  # noqa: E402

RESULT_DIR = OUT_DIR / "roberta_vs_llm"
SEED_RE = re.compile(r"^(?P<dataset>.+)-(?P<seed>\d+)$")
FILLS = {
    "green": "C6EFCE",
    "white": "FFFFFF",
    "yellow": "FFEB9C",
    "blue": "BDD7EE",
}


def roberta_predictions(train_path: Path, test_path: Path, args) -> pd.DataFrame:
    """roberta-large's test predictions for one pair; trains the pair first if they don't exist yet."""
    path = roberta_large.RESULT_DIR / f"{test_path.stem}.csv"
    if args.retrain or not path.exists():
        print(f"training roberta-large for {test_path.stem} ...")
        roberta_large.RESULT_DIR.mkdir(parents=True, exist_ok=True)
        roberta_large.run_pair(train_path, test_path, args)
    else:
        print(f"using existing roberta-large predictions: {path}")
    return pd.read_csv(path)


def compare_pair(train_path: Path, test_path: Path, args) -> pd.DataFrame:
    df = roberta_predictions(train_path, test_path, args)
    llm = llm_predictions(LLM_TEST_DIR / f"{test_path.stem}.csv")
    m = SEED_RE.match(test_path.stem)

    out = pd.DataFrame({
        "dataset": m["dataset"],
        "seed": m["seed"],
        "file": test_path.stem,
        TEXT_COLUMN: df[TEXT_COLUMN],
        "true": df["true"],
        "llm_pred": df[TEXT_COLUMN].map(llm),
        "roberta_pred": df["roberta_pred"],
    })
    out["llm_roberta_agree"] = out["llm_pred"] == out["roberta_pred"]
    out["llm_correct"] = out["llm_pred"] == out["true"]
    out["roberta_correct"] = out["roberta_pred"] == out["true"]
    out["agreement_color"] = out["llm_roberta_agree"].map({True: "green", False: "white"})
    out["correctness_color"] = [
        "blue" if llm_ok and rob_ok else "yellow" if rob_ok else "green" if llm_ok else "white"
        for llm_ok, rob_ok in zip(out["llm_correct"], out["roberta_correct"])
    ]
    return out


def summarize(table: pd.DataFrame) -> pd.DataFrame:
    """Share of sentences per category, per pair and (all seeds pooled) per dataset."""
    def shares(group: pd.DataFrame) -> pd.Series:
        return pd.Series({
            "rows": len(group),
            "agree": group["llm_roberta_agree"].mean(),
            "both_correct": (group["correctness_color"] == "blue").mean(),
            "only_roberta_correct": (group["correctness_color"] == "yellow").mean(),
            "only_llm_correct": (group["correctness_color"] == "green").mean(),
            "both_wrong": (group["correctness_color"] == "white").mean(),
            "llm_accuracy": group["llm_correct"].mean(),
            "roberta_accuracy": group["roberta_correct"].mean(),
        })

    per_pair = table.groupby(["dataset", "file"]).apply(shares, include_groups=False).reset_index()
    per_dataset = table.groupby("dataset").apply(shares, include_groups=False).reset_index()
    per_dataset["file"] = "all seeds"
    return pd.concat([per_pair, per_dataset], ignore_index=True).sort_values(["dataset", "file"])


def write_xlsx(table: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        table.to_excel(writer, index=False, sheet_name="comparison")
        sheet = writer.sheets["comparison"]
        fills = {color: PatternFill(start_color=hex_, end_color=hex_, fill_type="solid") for color, hex_ in FILLS.items()}
        for column in ("agreement_color", "correctness_color"):
            col_idx = table.columns.get_loc(column) + 1
            for row_idx, color in enumerate(table[column], start=2):
                sheet.cell(row=row_idx, column=col_idx).fill = fills[color]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--datasets", nargs="*", help="only datasets (test file name without seed) containing one of these")
    parser.add_argument("--retrain", action="store_true", help="train roberta-large again even if predictions exist")
    parser.add_argument("--trials", type=int, default=24,
                        help="Optuna trials when roberta-large has to be trained (16-point grid first, the rest is free)")
    parser.add_argument("--epochs", type=int, default=10, help="(maximum) epochs of roberta-large's default configuration")
    parser.add_argument("--retries", type=int, default=2, help="repeat a collapsed final training with another seed up to N times")
    parser.add_argument("--limit-train", type=int, help="subsample the train file to N rows (smoke test)")
    args = parser.parse_args()

    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    tables = []
    for test_path in sorted(TEST_DIR.glob("lab-manual-*-test-*.xlsx")):
        m = SEED_RE.match(test_path.stem)
        if not m or (args.datasets and not any(d in m["dataset"] for d in args.datasets)):
            continue
        train_path = paired_train(test_path)
        if not train_path.exists():
            print(f"skip {test_path.stem}: no train file {train_path.name}")
            continue
        tables.append(compare_pair(train_path, test_path, args))
    if not tables:
        raise SystemExit("no matching train/test pairs found")

    table = pd.concat(tables, ignore_index=True)
    table.to_csv(RESULT_DIR / "comparison.csv", index=False)
    write_xlsx(table, RESULT_DIR / "comparison.xlsx")
    summary = summarize(table)
    summary.to_csv(RESULT_DIR / "summary.csv", index=False)

    print("\n=== per dataset (all seeds pooled) ===")
    cols = ["dataset", "rows", "agree", "both_correct", "only_roberta_correct", "only_llm_correct", "both_wrong"]
    print(summary[summary["file"] == "all seeds"][cols].round(3).to_string(index=False))
    print(f"\nSaved: {RESULT_DIR / 'comparison.csv'}, {RESULT_DIR / 'comparison.xlsx'} and {RESULT_DIR / 'summary.csv'}")


if __name__ == "__main__":
    main()
